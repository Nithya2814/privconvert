import streamlit as st
import os
import io
import tempfile
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pdf2docx import Converter
from docx import Document
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Page configuration
st.set_page_config(
    page_title="PrivConvert - Local Document Converter",
    page_icon="🛡️",
    layout="centered"
)

st.title("🛡️ PrivConvert")
st.caption("100% Local & On-Device Document Converter | Zero Cloud Access")

# Helper: Apply PDF password encryption if requested
def apply_pdf_password(pdf_bytes: bytes, password: str) -> bytes:
    if not password:
        return pdf_bytes
    reader = PdfReader(io.BytesIO(pdf_bytes))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password=password, algorithm="AES-256")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()

# Tabs for direct practical conversions
tab1, tab2, tab3 = st.tabs([
    "📑 PDF to Word (.docx)", 
    "📝 Word / PPT to PDF", 
    "🖼️ Image to PDF (with Password Lock)"
])

# -------------------------------------------------------------
# TAB 1: PDF to Word
# -------------------------------------------------------------
with tab1:
    st.subheader("Convert PDF to Editable Word Document")
    pdf_file = st.file_uploader("Upload PDF file", type=["pdf"], key="pdf_to_word_uploader")
    
    if pdf_file:
        st.write(f"**File:** `{pdf_file.name}` ({round(pdf_file.size / 1024, 1)} KB)")
        if st.button("🔄 Convert to Word (.docx)", type="primary"):
            with st.spinner("Converting locally on your machine..."):
                try:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_in:
                        tmp_in.write(pdf_file.read())
                        tmp_in_path = tmp_in.name
                        
                    tmp_out_path = tmp_in_path.replace(".pdf", ".docx")
                    
                    # Convert using pdf2docx locally
                    cv = Converter(tmp_in_path)
                    cv.convert(tmp_out_path)
                    cv.close()
                    
                    with open(tmp_out_path, "rb") as f:
                        docx_bytes = f.read()
                        
                    # Cleanup temp files
                    os.remove(tmp_in_path)
                    os.remove(tmp_out_path)
                    
                    st.success("✅ Successfully converted to Word document!")
                    out_filename = f"{os.path.splitext(pdf_file.name)[0]}.docx"
                    st.download_button(
                        label=f"⬇️ Download {out_filename}",
                        data=docx_bytes,
                        file_name=out_filename,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                except Exception as e:
                    st.error(f"Conversion failed: {e}")

# -------------------------------------------------------------
# TAB 2: Word / PPT to PDF
# -------------------------------------------------------------
with tab2:
    st.subheader("Convert Word (.docx) or PowerPoint (.pptx) to PDF")
    doc_file = st.file_uploader("Upload Word or PowerPoint file", type=["docx", "pptx"], key="office_uploader")
    
    if doc_file:
        st.write(f"**File:** `{doc_file.name}` ({round(doc_file.size / 1024, 1)} KB)")
        
        # Optional Password protection
        lock_doc = st.checkbox("🔒 Protect Output PDF with Password", value=False, key="lock_office")
        doc_pass = ""
        if lock_doc:
            doc_pass = st.text_input("Enter password to lock PDF:", type="password", key="doc_p_in")
            
        if st.button("🚀 Convert to PDF", type="primary"):
            if lock_doc and not doc_pass:
                st.error("Please enter a password!")
            else:
                with st.spinner("Processing document locally..."):
                    try:
                        pdf_buffer = io.BytesIO()
                        doc_template = SimpleDocTemplate(pdf_buffer, pagesize=letter)
                        styles = getSampleStyleSheet()
                        story = []
                        
                        file_ext = os.path.splitext(doc_file.name)[1].lower()
                        
                        if file_ext == ".docx":
                            # Extract paragraphs from Word document
                            doc = Document(doc_file)
                            for para in doc.paragraphs:
                                text = para.text.strip()
                                if text:
                                    story.append(Paragraph(text, styles["Normal"]))
                                    story.append(Spacer(1, 8))
                                    
                        elif file_ext == ".pptx":
                            # Extract text from PowerPoint slides
                            prs = Presentation(doc_file)
                            for idx, slide in enumerate(prs.slides):
                                story.append(Paragraph(f"<b>Slide {idx + 1}</b>", styles["Heading2"]))
                                story.append(Spacer(1, 6))
                                for shape in slide.shapes:
                                    if shape.has_text_frame:
                                        for para in shape.text_frame.paragraphs:
                                            if para.text.strip():
                                                story.append(Paragraph(para.text.strip(), styles["Normal"]))
                                                story.append(Spacer(1, 4))
                                story.append(Spacer(1, 14))
                        
                        if not story:
                            story.append(Paragraph("Empty Document", styles["Normal"]))
                            
                        doc_template.build(story)
                        raw_pdf = pdf_buffer.getvalue()
                        
                        # Apply encryption if requested
                        final_pdf = apply_pdf_password(raw_pdf, doc_pass if lock_doc else None)
                        
                        st.success("✅ Converted to PDF successfully!")
                        out_pdf_name = f"{os.path.splitext(doc_file.name)[0]}.pdf"
                        st.download_button(
                            label=f"⬇️ Download {out_pdf_name}",
                            data=final_pdf,
                            file_name=out_pdf_name,
                            mime="application/pdf"
                        )
                    except Exception as e:
                        st.error(f"Error during conversion: {e}")

# -------------------------------------------------------------
# TAB 3: Image to PDF (with Password Lock)
# -------------------------------------------------------------
with tab3:
    st.subheader("Convert Images to Password-Protected PDF")
    img_file = st.file_uploader("Upload Image (PNG, JPG, WEBP)", type=["png", "jpg", "jpeg", "webp"], key="img_pdf_uploader")
    
    if img_file:
        st.image(img_file, width=300)
        
        lock_img = st.checkbox("🔒 Protect PDF with Password", value=True, key="img_lock_chk")
        img_pass = ""
        if lock_img:
            img_pass = st.text_input("Enter secret password:", type="password", key="img_p_in")
            
        if st.button("🚀 Convert Image to PDF", type="primary"):
            if lock_img and not img_pass:
                st.error("Please enter a password!")
            else:
                with st.spinner("Converting locally in RAM..."):
                    try:
                        im = Image.open(img_file)
                        if im.mode in ("RGBA", "P"):
                            im = im.convert("RGB")
                        
                        buf = io.BytesIO()
                        im.save(buf, format="PDF")
                        raw_pdf = buf.getvalue()
                        
                        final_pdf = apply_pdf_password(raw_pdf, img_pass if lock_img else None)
                        
                        st.success("✅ Image converted to PDF!")
                        out_img_pdf = f"{os.path.splitext(img_file.name)[0]}.pdf"
                        st.download_button(
                            label=f"⬇️ Download {out_img_pdf}",
                            data=final_pdf,
                            file_name=out_img_pdf,
                            mime="application/pdf"
                        )
                    except Exception as e:
                        st.error(f"Error: {e}")
